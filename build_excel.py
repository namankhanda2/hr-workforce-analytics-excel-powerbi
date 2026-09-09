"""
Builds the Excel workbook: HR_Workforce_Analytics.xlsx

Workbook sheets:
  Guide          - overview + methodology + navigation
  Raw Data       - messy source data (quality issues flagged, for cleaning demo)
  Clean Data     - cleaned + derived columns (as an Excel Table for Power BI import)
  Insights       - formula-driven pivot-style summaries + native Excel charts
  Dashboard      - KPI cards + charts + key findings

Run:  .venv/bin/python build_excel.py
"""

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent
CLEAN_CSV = BASE / "data" / "hr_cleaned.csv"
RAW_CSV = BASE / "data" / "raw" / "hr_employees_raw.csv"
OUT = BASE / "excel" / "HR_Workforce_Analytics.xlsx"

N_ROWS = 3000
LAST = N_ROWS + 1  # last data row in sheets (row 1 = header)

DARK = "1F3864"
MID = "2E75B6"
ACCENT = "548235"
RED = "C00000"
WHITE = "FFFFFF"

BORDER_BOX = "thin"
HDR_FILL = PatternFill("solid", fgColor=DARK)
BORDER = Border(
    left=Side(style=BORDER_BOX, color="BFBFBF"),
    right=Side(style=BORDER_BOX, color="BFBFBF"),
    top=Side(style=BORDER_BOX, color="BFBFBF"),
    bottom=Side(style=BORDER_BOX, color="BFBFBF"),
)

DEPARTMENTS = [
    "Executive", "Finance", "Human Resources", "IT", "Operations",
    "Marketing", "Sales", "Customer Support", "Research & Development",
]
EXIT_REASONS = [
    "Resigned", "Retirement", "Performance / Termination",
    "Layoff", "Relocation",
]
YEARS = [2020, 2021, 2022, 2023, 2024, 2025]

CLEAN_COLS = [
    "EmployeeID", "Gender", "Age", "MaritalStatus", "Education",
    "Department", "JobRole", "JobLevel", "OfficeLocation",
    "EmploymentType", "HireDate", "ExitDate", "IsActive", "ExitReason",
    "SalaryAnnual", "BonusPct", "PerformanceRating", "SatisfactionScore",
    "TenureMonths", "AbsenteeDays",
]
CLEAN_LETTER = {name: get_column_letter(i + 1) for i, name in enumerate(CLEAN_COLS)}

DEPT_C = f"'Clean Data'!${CLEAN_LETTER['Department']}$2:${CLEAN_LETTER['Department']}${LAST}"
SAL_C = f"'Clean Data'!${CLEAN_LETTER['SalaryAnnual']}$2:${CLEAN_LETTER['SalaryAnnual']}${LAST}"
LVL_C = f"'Clean Data'!${CLEAN_LETTER['JobLevel']}$2:${CLEAN_LETTER['JobLevel']}${LAST}"
GEN_C = f"'Clean Data'!${CLEAN_LETTER['Gender']}$2:${CLEAN_LETTER['Gender']}${LAST}"
ACT_C = f"'Clean Data'!${CLEAN_LETTER['IsActive']}$2:${CLEAN_LETTER['IsActive']}${LAST}"
HIRE_C = f"'Clean Data'!${CLEAN_LETTER['HireDate']}$2:${CLEAN_LETTER['HireDate']}${LAST}"
REASON_C = f"'Clean Data'!${CLEAN_LETTER['ExitReason']}$2:${CLEAN_LETTER['ExitReason']}${LAST}"
AGE_C = f"'Clean Data'!${CLEAN_LETTER['Age']}$2:${CLEAN_LETTER['Age']}${LAST}"
ID_C = f"'Clean Data'!${CLEAN_LETTER['EmployeeID']}$2:${CLEAN_LETTER['EmployeeID']}${LAST}"


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def fmt_yesno(v):
    return "Yes" if v in ("True", "Yes", "1") else "No"


def parseno(v):
    if not v or v == "":
        return ""
    return "Yes" if v in ("True", "Yes") else "No"


def header_row(ws, row, labels):
    for i, lab in enumerate(labels):
        cell = ws.cell(row=row, column=1 + i, value=lab)
        cell.font = Font(bold=True, color=WHITE, size=11)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


# ---------------------------------------------------------------- sheets
def build_guide(ws, raw_rows, clean_rows, active):
    ws.sheet_properties.tabColor = MID
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100

    ws["B1"] = "HR Workforce Analytics - Excel + Power BI Project"
    ws["B1"].font = Font(bold=True, size=18, color=DARK)

    rows = [
        ("What is this?", "Synthetic HR data for 3,000 employees analysed end-to-end in Excel, then handed to Power BI."),
        ("Sheets", "Raw Data -> Clean Data -> Insights -> Dashboard. Review in that order."),
        ("Raw Data", "As-collected file: duplicates, blank salary/department, gender typos, salary outliers, age typos."),
        ("Clean Data", "Quality fixes applied; derived columns use live Excel formulas (IFS). Exported as Excel Table tCleanData for one-click Power BI import."),
        ("Insights", "Pivot-style summaries built from SUMIFS / COUNTIFS / AVERAGEIFS + native Excel charts."),
        ("Dashboard", "KPI cards, charts and findings for the HR manager."),
        ("Power BI", "Power BI-ready star schema (DimEmployee, DimDate, FactEmploymentEvents, FactCompensation) sits in data/powerbi/. See powerbi/PowerBI_Build_Guide.md."),
        ("", ""),
        ("Cleaning log", f"{len(raw_rows)} raw rows scanned -> {len(clean_rows)} clean rows. What was fixed:"),
    ]
    r = 2
    for label, txt in rows:
        if label:
            ws.cell(row=r, column=2, value=label).font = Font(bold=True, color=DARK)
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=3)
            ws.cell(row=r, column=3, value=txt).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    fixes = [
        "10 duplicate employee rows removed",
        "blank SalaryAnnual rows removed (unusable for compensation analytics)",
        "blank Department rows removed",
        "'Unknown' + lower-case Gender values normalised; bad values removed",
        "salary outliers (>$500k) excluded",
        "Age typos (200+) excluded",
        "derived columns AgeBand, SalaryBand, AttritionStatus added (live Excel formulas)",
    ]
    for f_ in fixes:
        ws.cell(row=r, column=3, value="* " + f_)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Headline numbers").font = Font(bold=True, color=DARK)
    r += 1
    for label, value in [
        ("Total employees", len(clean_rows)),
        ("Active", active),
        ("Attrition rate", f"{100 * (1 - active / len(clean_rows)):.1f}%"),
    ]:
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=value)
        r += 1


def build_raw(ws, raw_rows, clean_rows):
    ws.sheet_properties.tabColor = RED

    clean_ids = {r["EmployeeID"] for r in clean_rows}
    raw_ids = {r["EmployeeID"] for r in raw_rows}

    headers = CLEAN_COLS + ["RowStatus"]
    header_row(ws, 1, headers)
    ncols = len(headers)

    seen = set()
    r = 2
    for row in raw_rows:
        flag = []
        eid = row["EmployeeID"]
        if eid in seen:
            flag.append("Duplicate")
        elif eid not in clean_ids:
            flag.append("Removed in cleaning")
        else:
            if not row.get("SalaryAnnual"):
                flag.append("Blank salary")
            if row["SalaryAnnual"] in ("999999", "1000000", "1500000"):
                flag.append("Outlier salary")
            if (row.get("Gender") or "").lower() not in ("male", "female"):
                flag.append("Bad gender")
            if not row.get("Department"):
                flag.append("Blank department")
            if row["Age"] in ("200", "210", "250"):
                flag.append("Age typo")
        seen.add(eid)

        vals = []
        for col in CLEAN_COLS:
            v = row.get(col, "") or ""
            if col in ("HireDate", "ExitDate") and v:
                v = date.fromisoformat(v[:10])
            elif col == "IsActive":
                v = fmt_yesno(v)
            vals.append(v)
        vals.append("; ".join(flag) if flag else "OK")
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = BORDER
            if col_numfmt := {
                3: "0", 11: "yyyy-mm-dd", 12: "yyyy-mm-dd", 15: "$#,##0",
            }.get(j):
                cell.number_format = col_numfmt
        r += 1

    last_letter = get_column_letter(ncols)
    ws.auto_filter.ref = f"A1:{last_letter}{r - 1}"
    ws.freeze_panes = "A2"

    ws.conditional_formatting.add(
        f"O2:O{r - 1}",
        ColorScaleRule(start_type="min", start_color="C6EFCE",
                       end_type="max", end_color="FFC7CE"),
    )
    ws.conditional_formatting.add(
        f"{last_letter}2:{last_letter}{r - 1}",
        CellIsRule(operator="notEqual", formula=['"OK"'],
                   fill=PatternFill("solid", fgColor="FFC7CE")),
    )
    ws["U1"] = "Quality Flag"
    for col, w in {"A": 12, "B": 10, "C": 7, "K": 12, "L": 12, "M": 10,
                   "N": 25, "O": 13, "P": 10, "U": 22}.items():
        ws.column_dimensions[col].width = w


def build_clean(ws):
    ws.sheet_properties.tabColor = ACCENT
    headers = CLEAN_COLS + ["AgeBand", "SalaryBand", "AttritionStatus"]
    header_row(ws, 1, headers)

    for i, row in enumerate(read_csv(CLEAN_CSV), start=2):
        vals = []
        for col in CLEAN_COLS:
            v = row.get(col, "") or ""
            if col in ("HireDate", "ExitDate") and v:
                v = date.fromisoformat(v[:10])
            elif col == "IsActive":
                v = parseno(v)
            vals.append(v)
        age, sal, act = (f"{CLEAN_LETTER['Age']}{i}",
                         f"{CLEAN_LETTER['SalaryAnnual']}{i}",
                         f"{CLEAN_LETTER['IsActive']}{i}")
        vals += [
            f'=IFS({age}<30,"Under 30",{age}<40,"30-39",{age}<50,"40-49",TRUE,"50+")',
            f'=IFS({sal}<60000,"Low (<60k)",{sal}<100000,"Mid (60-100k)",TRUE,"High (100k+)")',
            f'=IF({act}="No","Attrited","Active")',
        ]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = BORDER
            if nf := {3: "0", 11: "yyyy-mm-dd", 12: "yyyy-mm-dd", 15: "$#,##0"}.get(j):
                cell.number_format = nf

    ncols = len(headers)
    last_letter = get_column_letter(ncols)
    ws.auto_filter.ref = f"A1:{last_letter}{LAST}"
    ws.freeze_panes = "A2"

    tab = Table(displayName="tCleanData", ref=f"A1:{last_letter}{LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)

    for col, w in {"A": 12, "B": 10, "C": 7, "D": 12, "E": 12, "F": 22, "G": 20,
                   "H": 9, "I": 13, "J": 14, "K": 12, "L": 12, "M": 10, "N": 25,
                   "O": 13, "P": 9, "Q": 10, "R": 12, "S": 11, "T": 11,
                   "U": 12, "V": 15, "W": 13}.items():
        ws.column_dimensions[col].width = w


def write_body(ws, r, start_col, end_col, nfmts=None):
    for c in range(start_col, end_col + 1):
        ws.cell(row=r, column=c).border = BORDER
    for c, fmt in (nfmts or {}).items():
        ws.cell(row=r, column=c).number_format = fmt


def build_insights(ws):
    ws.sheet_properties.tabColor = MID

    def sec(r, c, text):
        ws.cell(row=r, column=c, value=text).font = Font(bold=True, size=12, color=DARK)

    # ---- 1) Headcount & attrition by dept
    sec(1, 1, "1) Headcount and attrition rate by department")
    header_row(ws, 2, ["Department", "Headcount", "Avg Salary", "Attrited", "Attrition Rate"])
    r0 = 3
    for i, dept in enumerate(DEPARTMENTS):
        r = r0 + i
        ws.cell(row=r, column=1, value=dept)
        ws.cell(row=r, column=2, value=f"=COUNTIF({DEPT_C},A{r})")
        ws.cell(row=r, column=3, value=f"=IF(B{r}=0,0,ROUND(AVERAGEIFS({SAL_C},{DEPT_C},A{r}),0))")
        ws.cell(row=r, column=4, value=f'=COUNTIFS({DEPT_C},A{r},{ACT_C},"No")')
        ws.cell(row=r, column=5, value=f"=IF(B{r}=0,0,D{r}/B{r})")
        write_body(ws, r, 1, 5, nfmts={3: "$#,##0", 5: "0.0%"})
    tr = r0 + len(DEPARTMENTS)
    ws.cell(row=tr, column=1, value="Total").font = Font(bold=True, color=WHITE)
    ws.cell(row=tr, column=2, value=f"=SUM(B{r0}:B{tr-1})")
    ws.cell(row=tr, column=3, value=f"=ROUND(AVERAGE({SAL_C}),0)")
    ws.cell(row=tr, column=4, value=f"=SUM(D{r0}:D{tr-1})")
    ws.cell(row=tr, column=5, value=f"=IF(B{tr}=0,0,D{tr}/B{tr})")
    for c in range(1, 6):
        cell = ws.cell(row=tr, column=c)
        cell.fill = HDR_FILL
    write_body(ws, tr, 1, 5, nfmts={3: "$#,##0", 5: "0.0%"})

    ws.conditional_formatting.add(
        f"E{r0}:E{tr-1}",
        ColorScaleRule(start_type="num", start_value=0, start_color="C6EFCE",
                       end_type="num", end_value=0.5, end_color="FFC7CE"),
    )

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Headcount by Department"
    bar.height, bar.width = 9, 16
    bar.legend = None
    bar.add_data(Reference(ws, min_col=2, min_row=2, max_row=tr - 1), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=tr - 1))
    ws.add_chart(bar, "G2")

    # ---- 2) Gender pay gap by level
    sec(14, 1, "2) Average salary by gender and job level (pay-gap snapshot)")
    header_row(ws, 15, ["Job Level", "Female", "Male", "Pay Gap (M vs F)"])
    for lvl in range(1, 8):
        r = 15 + lvl
        ws.cell(row=r, column=1, value=f"Level {lvl}")
        ws.cell(row=r, column=2, value=f'=ROUND(AVERAGEIFS({SAL_C},{LVL_C},"Level {lvl}",{GEN_C},"Female"),0)')
        ws.cell(row=r, column=3, value=f'=ROUND(AVERAGEIFS({SAL_C},{LVL_C},"Level {lvl}",{GEN_C},"Male"),0)')
        ws.cell(row=r, column=4, value=f"=IF(B{r}=0,0,(C{r}-B{r})/B{r})")
        write_body(ws, r, 1, 4, nfmts={2: "$#,##0", 3: "$#,##0", 4: "0.0%"})

    # ---- 3) Hires per year
    sec(25, 1, "3) Hiring trend by year")
    header_row(ws, 26, ["Year", "Hires"])
    for i, y in enumerate(YEARS):
        r = 27 + i
        ws.cell(row=r, column=1, value=y)
        ws.cell(row=r, column=2, value=f'=COUNTIFS({HIRE_C},">="&DATE({y},1,1),{HIRE_C},"<="&DATE({y},12,31))')
        write_body(ws, r, 1, 2)

    line = LineChart()
    line.title = "Hires by Year"
    line.height, line.width = 9, 16
    line.add_data(Reference(ws, min_col=2, min_row=26, max_row=27 + len(YEARS) - 1), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=27, max_row=27 + len(YEARS) - 1))
    ws.add_chart(line, "G16")

    # ---- 4) Exits by reason
    sec(36, 1, "4) Departures by exit reason")
    header_row(ws, 37, ["Exit Reason", "Count"])
    for i, reason in enumerate(EXIT_REASONS):
        r = 38 + i
        ws.cell(row=r, column=1, value=reason)
        ws.cell(row=r, column=2, value=f"=COUNTIF({REASON_C},A{r})")
        write_body(ws, r, 1, 2)

    pie = PieChart()
    pie.title = "Departures by Exit Reason"
    pie.height, pie.width = 9, 14
    pie.add_data(Reference(ws, min_col=2, min_row=37, max_row=37 + len(EXIT_REASONS)), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=38, max_row=38 + len(EXIT_REASONS) - 1))
    ws.add_chart(pie, "G36")

    for col, w in {"A": 27, "B": 11, "C": 13, "D": 11, "E": 14, "F": 3}.items():
        ws.column_dimensions[col].width = w


def build_dashboard(ws, ins, clean_rows, active):
    ws.sheet_properties.tabColor = DARK

    ws.merge_cells("A1:J1")
    c = ws["A1"]
    c.value = "HR WORKFORCE DASHBOARD"
    c.font = Font(bold=True, size=20, color=WHITE)
    c.fill = PatternFill("solid", fgColor=DARK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:J2")
    c2 = ws["A2"]
    c2.value = "Snapshot: 31-Dec-2025   |   Synthetic HR data   |   Built in Excel, Power BI-ready"
    c2.font = Font(italic=True, color=MID)
    c2.alignment = Alignment(horizontal="center")

    # KPI cards (2 columns each)
    cards = [
        ("TOTAL HEADCOUNT", f"=COUNTA({ID_C})", "#,##0", DARK),
        ("ACTIVE EMPLOYEES", f'=COUNTIF({ACT_C},"Yes")', "#,##0", MID),
        ("ATTRITION RATE", f'=1-COUNTIF({ACT_C},"Yes")/COUNTA({ID_C})', "0.0%", RED),
        ("AVG SALARY (USD)", f"=ROUND(AVERAGE({SAL_C}),0)", "$#,##0", ACCENT),
        ("AVG TENURE (MONTHS)", f"=ROUND(AVERAGE('Clean Data'!$S$2:$S${LAST}),1)", "0.0", MID),
        ("AVG SATISFACTION (/10)", f"=ROUND(AVERAGE('Clean Data'!$R$2:$R${LAST}),1)", "0.0", ACCENT),
    ]
    for i, (label, formula, fmt, color) in enumerate(cards):
        c1 = 1 + i * 2
        row_label, row_val = 4, 5
        ws.merge_cells(start_row=row_label, start_column=c1, end_row=row_label, end_column=c1 + 1)
        ws.merge_cells(start_row=row_val, start_column=c1, end_row=row_val, end_column=c1 + 1)
        lc = ws.cell(row=row_label, column=c1, value=label)
        lc.font = Font(bold=True, size=10, color=WHITE)
        lc.fill = PatternFill("solid", fgColor=color)
        lc.alignment = Alignment(horizontal="center", vertical="center")
        vc = ws.cell(row=row_val, column=c1, value=formula)
        vc.number_format = fmt
        vc.font = Font(bold=True, size=16, color=DARK)
        vc.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row_label].height = 22
        ws.row_dimensions[row_val].height = 30

    ws.cell(row=8, column=1, value="Workforce snapshots (live formulas, see Insights)").font = Font(bold=True, color=DARK)

    # Charts sourced from Insights sheet (cross-sheet references)
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Headcount by Department"
    bar.height, bar.width = 9, 16
    bar.legend = None
    bar.add_data(Reference(ins, min_col=2, min_row=2, max_row=11), titles_from_data=True)
    bar.set_categories(Reference(ins, min_col=1, min_row=3, max_row=11))
    ws.add_chart(bar, "A10")

    line = LineChart()
    line.title = "Hires by Year"
    line.height, line.width = 9, 15
    line.add_data(Reference(ins, min_col=2, min_row=26, max_row=32), titles_from_data=True)
    line.set_categories(Reference(ins, min_col=1, min_row=27, max_row=32))
    ws.add_chart(line, "H10")

    # Findings
    ws.cell(row=22, column=1, value="KEY FINDINGS & RECOMMENDATIONS").font = Font(bold=True, size=12, color=DARK)
    attrition = 1 - active / len(clean_rows)
    findings = [
        f"Headcount {len(clean_rows)} (active {active}, attrition {attrition:.1%}).",
        "Attrition skews to Sales, Customer Support and Operations - review comp and workload there.",
        "Exits concentrate in short tenures (<18 months) and junior levels - strengthen onboarding.",
        "Run the pay-gap table (Insights sheet) each quarter; flag any level where gap > 10%.",
        "Use exit-reason trend (Insights) to target retention: e.g. offer flexibility to reduce 'Relocation'.",
    ]
    r = 23
    for f_ in findings:
        ws.cell(row=r, column=1, value="* " + f_)
        r += 1

    for col, w in {"A": 14, "B": 14, "C": 14, "D": 14, "E": 14,
                   "F": 14, "G": 14, "H": 14, "I": 14, "J": 14}.items():
        ws.column_dimensions[col].width = w


def main():
    raw_rows = read_csv(RAW_CSV)
    clean_rows = read_csv(CLEAN_CSV)
    active = sum(1 for r in clean_rows if r["IsActive"] in ("True", "Yes"))

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True

    guide = wb.create_sheet("Guide")
    raw = wb.create_sheet("Raw Data")
    cln = wb.create_sheet("Clean Data")
    ins = wb.create_sheet("Insights")
    dsh = wb.create_sheet("Dashboard")

    build_guide(guide, raw_rows, clean_rows, active)
    build_raw(raw, raw_rows, clean_rows)
    build_clean(cln)
    build_insights(ins)
    build_dashboard(dsh, ins, clean_rows, active)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Workbook written -> {OUT}")


if __name__ == "__main__":
    main()